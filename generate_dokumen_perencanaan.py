import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_planning_document():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    font_family = "Arial"
    
    font_title = Font(name=font_family, size=11, bold=True, color="000000")
    font_header = Font(name=font_family, size=10, bold=True, color="000000")
    font_bold = Font(name=font_family, size=10, bold=True, color="000000")
    font_regular = Font(name=font_family, size=10, bold=False, color="000000")
    font_table_name = Font(name=font_family, size=10, bold=True, color="000000")

    thin_border_side = Side(border_style="thin", color="D3D3D3")
    full_thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    fill_cyan = PatternFill(start_color="00FFFF", end_color="00FFFF", fill_type="solid")
    fill_schema_header = PatternFill(start_color="B6D7A8", end_color="B6D7A8", fill_type="solid")
    fill_schema_table = PatternFill(start_color="9FC5E8", end_color="9FC5E8", fill_type="solid")

    palette = [
        "FFEAD1DC",  # Soft pinkish lavender
        "FFD9D2E9",  # Soft lavender
        "FFCFE2F3",  # Soft blue
        "FFFFF2CC",  # Soft warm yellow (cream)
        "FFF4CCCC",  # Soft light coral/peach
        "FFD9EAD3",  # Soft green
        "FFFCE5CD",  # Soft apricot
    ]

    # =========================================================================
    # 1. SHEET: Deskripsi
    # =========================================================================
    ws_desc = wb.create_sheet(title="Deskripsi")
    ws_desc.views.sheetView[0].showGridLines = True

    desc_rows = [
        ("DESKRIPSI", "Aplikasi Deteksi dan Verifikasi Berita Hoaks Menggunakan Machine Learning dan NLP Berbasis Mobile & Web"),
        ("Tujuan :", (
            "1. Mengembangkan aplikasi mobile berbasis Flutter (HoaxCheck) yang intuitif dan responsif untuk membantu masyarakat mendeteksi serta memverifikasi potensi berita hoaks atau pesan berantai secara instan.\n"
            "2. Menerapkan algoritma Machine Learning dan pemodelan Natural Language Processing (NLP) seperti clustering teks (K-Means, DBSCAN, HDBSCAN, Hierarchical) serta ekstraksi fitur linguistik (TF-IDF, IndoBERT) untuk mengelompokkan topik disinformasi dan mendeteksi probabilitas keabsahan informasi.\n"
            "3. Mengembangkan sistem deteksi pola kebahasaan manipulatif pada teks berita (bahasa provokatif, ajakan mendesak segera menyebarkan, ketiadaan rujukan media resmi) serta menyajikan visualisasi tingkat keyakinan (Confidence Score).\n"
            "4. Mengembangkan dashboard Web Admin berbasis Flask yang dilengkapi otomatisasi RSS Scraper berita valid dari portal media kredibel (Antara News, CNN Indonesia, CNBC Indonesia), pengelolaan korpus dataset hoaks Komdigi, dan monitoring metrik model.\n"
            "5. Mengintegrasikan ekosistem Web Admin, Machine Learning inference engine, RESTful Cloud API, dan aplikasi mobile Flutter secara komprehensif, aman, dan berkesinambungan."
        )),
        ("Latar Belakang : ", (
            "Penyebaran hoaks, disinformasi, dan misinformasi telah menjadi ancaman serius dalam lanskap informasi digital di Indonesia, terutama yang beredar masif di aplikasi pesan instan (grup WhatsApp, Telegram) dan media sosial. Sering kali masyarakat awam kesulitan mengenali apakah suatu informasi faktual atau hasil manipulasi, sehingga rentan memicu kepanikan massal, kerugian finansial, hingga perpecahan sosial. Sebagian besar pengecekan fakta yang ada saat ini masih bergantung pada verifikasi manual yang memakan waktu lama.\n\n"
            "Oleh karena itu, dikembangkan sistem HoaxCheck, sebuah platform cerdas verifikasi berita berbasis Machine Learning dan NLP. HoaxCheck menghadirkan aplikasi mobile yang dilengkapi fitur deteksi cepat dari papan klip (auto-detect clipboard), kalkulasi confidence score cerdas, identifikasi pola indikator kebahasaan manipulatif, serta penelusuran riwayat verifikasi yang terhubung ke cloud server dan web admin pengelola korpus berita valid dan hoaks Komdigi."
        )),
        ("Pengguna : \n", (
            "Pengguna\n"
            "1. Admin - Mengelola dataset hoaks dan berita valid, menjalankan scraping berita terkini, mengevaluasi metrik klastering/model NLP, dan memantau log inferensi melalui Web Admin.\n"
            "2. Pengguna umum - Menggunakan aplikasi mobile HoaxCheck untuk memeriksa keaslian teks klaim berita, meninjau rincian analisis linguistik & confidence score, mengelola riwayat verifikasi, serta membaca panduan edukasi literasi digital."
        )),
        ("Lingkup Sistem :", (
            "Web Admin: Mengelola berkas dataset (.csv, .xlsx), scraper RSS berita valid multi-media, monitoring metrik model klastering, dan audit log sistem.\n"
            "Machine Learning: Pemrosesan awal teks (cleaning, tokenizing, stopword, vectorizing), clustering topik isu hoaks (K-Means, DBSCAN, HDBSCAN, Hierarchical), klasifikasi probabilitas hoaks vs valid, serta evaluasi metrik (Silhouette Score, Davies-Bouldin, Calinski-Harabasz).\n"
            "Mobile Flutter: Tampilan beranda interaktif, deteksi hoaks instan dengan fitur auto-detect clipboard, kartu rincian analisis NLP (confidence circular dial & indikator manipulatif), riwayat pemeriksaan dengan filter status, serta pengaturan preferensi profil (termasuk Dark Mode).\n"
            "Cloud/API: Endpoint inferensi teks real-time, validasi skema input, pengiriman data statistik agregasi hoaks, serta manajemen CORS dan keamanan payload."
        )),
        ("Teknologi Pemrograman : ", "Python, Flask, Flutter (Dart), RESTful API, Scikit-Learn, HDBSCAN, IndoBERT / TF-IDF, Pandas, SQLite / PostgreSQL, XML Parser (RSS)")
    ]

    ws_desc.column_dimensions["A"].width = 28.0
    ws_desc.column_dimensions["B"].width = 135.0

    row_heights_desc = {1: 26.0, 2: 95.0, 3: 95.0, 4: 55.0, 5: 75.0, 6: 28.0}

    for r_idx, (label, val) in enumerate(desc_rows, start=1):
        cell_a = ws_desc.cell(row=r_idx, column=1, value=label)
        cell_b = ws_desc.cell(row=r_idx, column=2, value=val)

        cell_a.font = font_bold
        cell_a.fill = fill_cyan
        cell_a.alignment = Alignment(vertical="center", wrap_text=True)
        cell_a.border = full_thin_border

        cell_b.font = font_regular
        cell_b.alignment = Alignment(vertical="center", wrap_text=True)
        cell_b.border = full_thin_border

        if r_idx in row_heights_desc:
            ws_desc.row_dimensions[r_idx].height = row_heights_desc[r_idx]

    # =========================================================================
    # Helper to populate Feature Sheets (Mobile, Web, Cloud)
    # =========================================================================
    def populate_feature_sheet(ws, title_text, feature_groups):
        ws.views.sheetView[0].showGridLines = True
        
        # Column widths matching House Price template
        ws.column_dimensions["A"].width = 7.5
        ws.column_dimensions["B"].width = 22.0
        ws.column_dimensions["C"].width = 28.0
        ws.column_dimensions["D"].width = 17.5
        ws.column_dimensions["E"].width = 75.0
        ws.column_dimensions["F"].width = 58.0

        # Row 1: FITUR title
        ws.row_dimensions[1].height = 20.0
        ws.row_dimensions[2].height = 24.0

        cell_title = ws.cell(row=1, column=1, value=title_text)
        cell_title.font = font_title
        cell_title.fill = fill_cyan
        cell_title.alignment = Alignment(horizontal="center", vertical="center")
        cell_title.border = full_thin_border

        for c in range(2, 7):
            c_cell = ws.cell(row=1, column=c, value="")
            c_cell.border = full_thin_border

        # Row 2: Headers
        headers = ["No", "Fitur", "Sub Fitur", "Hak Akses", "Deskripsi Fitur", "Inputan"]
        for c_idx, h_text in enumerate(headers, start=1):
            c_cell = ws.cell(row=2, column=c_idx, value=h_text)
            c_cell.font = font_header
            c_cell.fill = fill_cyan
            c_cell.alignment = Alignment(horizontal="center", vertical="center")
            c_cell.border = full_thin_border

        current_row = 3
        for g_idx, group in enumerate(feature_groups):
            color_hex = palette[g_idx % len(palette)]
            group_fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")

            start_row = current_row
            num_sub = len(group["subfeatures"])
            end_row = start_row + num_sub - 1

            for s_idx, sf in enumerate(group["subfeatures"]):
                r = start_row + s_idx
                ws.row_dimensions[r].height = 32.0

                c1 = ws.cell(row=r, column=1)
                c2 = ws.cell(row=r, column=2)
                c3 = ws.cell(row=r, column=3, value=sf["name"])
                c4 = ws.cell(row=r, column=4, value=sf["access"])
                c5 = ws.cell(row=r, column=5, value=sf["desc"])
                c6 = ws.cell(row=r, column=6, value=sf["input"])

                if s_idx == 0:
                    c1.value = float(g_idx + 1)
                    c2.value = group["name"]

                c1.font = font_regular
                c2.font = font_regular
                c3.font = font_regular
                c4.font = font_regular
                c5.font = font_regular
                c6.font = font_regular

                c1.alignment = Alignment(horizontal="center", vertical="center")
                c2.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                c4.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c5.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                c6.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

                for c in [c1, c2, c3, c4, c5, c6]:
                    c.fill = group_fill
                    c.border = full_thin_border

            # Merge columns A and B vertically if multiple subfeatures
            if end_row > start_row:
                ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
                ws.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)

            current_row = end_row + 1

    # =========================================================================
    # 2. SHEET: Fitur - Mobile
    # =========================================================================
    ws_mobile = wb.create_sheet(title="Fitur - Mobile")
    mobile_groups = [
        {
            "name": "Beranda (Home)",
            "subfeatures": [
                {
                    "name": "Tampilan ringkasan beranda",
                    "access": "Pengguna",
                    "desc": "Menampilkan salam pembuka (greeting pengguna), visualisasi ringkasan aktivitas verifikasi informasi, dan edukasi literasi digital di beranda utama.",
                    "input": "-"
                },
                {
                    "name": "Banner deteksi cepat cek hoaks",
                    "access": "Pengguna",
                    "desc": "Menyajikan kartu ajakan interaktif (hero banner) untuk mengarahkan pengguna secara instan ke halaman pengujian teks berita/klaim.",
                    "input": "Tombol 'Mulai Periksa'"
                },
                {
                    "name": "Aksi cepat (Quick Actions)",
                    "access": "Pengguna",
                    "desc": "Menyediakan jalan pintas untuk tempel teks instan dari papan klip, menulis klaim manual, panduan cek fakta, dan tips literasi digital.",
                    "input": "Pilihan kartu aksi cepat"
                },
                {
                    "name": "Statistik pemeriksaan pengguna",
                    "access": "Pengguna",
                    "desc": "Menyajikan metrik ringkasan aktivitas pemeriksaan 30 hari terakhir (Total klaim diuji, persentase terindikasi hoaks, dan berita valid).",
                    "input": "-"
                },
                {
                    "name": "Log pemeriksaan terakhir (Recent Check)",
                    "access": "Pengguna",
                    "desc": "Menampilkan kartu pratinjau hasil analisis teks yang terakhir kali diuji beserta label status, waktu uji WIB, dan sumber masukan.",
                    "input": "Tombol 'Detail Analisis'"
                },
                {
                    "name": "Tips & edukasi mengenali hoaks",
                    "access": "Pengguna",
                    "desc": "Menampilkan galeri horizontal (korsel) tips praktis mengenali ciri-ciri hoaks dan panduan verifikasi ke domain resmi pemerintah/media terpercaya.",
                    "input": "Geser kartu tips (swipe)"
                }
            ]
        },
        {
            "name": "Deteksi Teks Hoaks (Fitur NLP & ML)",
            "subfeatures": [
                {
                    "name": "Deteksi otomatis clipboard",
                    "access": "Pengguna",
                    "desc": "Mendeteksi keberadaan teks pada papan klip perangkat secara otomatis saat halaman dibuka dan menampilkan banner 'Gunakan' untuk menempel langsung.",
                    "input": "Aksi 'Gunakan' / 'Tutup'"
                },
                {
                    "name": "Form input teks klaim berita",
                    "access": "Pengguna",
                    "desc": "Menginputkan narasi berita, pesan berantai WhatsApp, atau klaim yang ingin diverifikasi kebenarannya ke dalam area teks multi-baris.",
                    "input": "Teks klaim berita (Textfield multi-line)"
                },
                {
                    "name": "Pembersih teks instan",
                    "access": "Pengguna",
                    "desc": "Mengosongkan seluruh isi teks masukan pada form input pemeriksaan klaim hanya dengan satu kali ketukan tombol.",
                    "input": "Tombol 'Bersihkan'"
                },
                {
                    "name": "Eksekusi analisis inferensi NLP",
                    "access": "Pengguna",
                    "desc": "Mengirimkan teks masukan ke pipeline inferensi model machine learning untuk diproses dan dianalisis secara cerdas.",
                    "input": "Tombol 'Periksa Keaslian Teks'"
                },
                {
                    "name": "Indikator progres & estimasi waktu",
                    "access": "Pengguna",
                    "desc": "Menampilkan animasi status pemrosesan analisis teks serta durasi latensi komputasi inferensi (dalam satuan detik).",
                    "input": "-"
                }
            ]
        },
        {
            "name": "Rincian Hasil Analisis (Detail Analisis)",
            "subfeatures": [
                {
                    "name": "Visualisasi skor keyakinan (Confidence Dial)",
                    "access": "Pengguna",
                    "desc": "Menampilkan persentase tingkat keyakinan model inferensi dalam bentuk Circular Progress Indicator dan badge status (Terindikasi Hoaks / Valid).",
                    "input": "-"
                },
                {
                    "name": "Identifikasi indikator linguistik manipulatif",
                    "access": "Pengguna",
                    "desc": "Menyajikan rincian kartu pola kebahasaan manipulatif yang terdeteksi (bahasa sensasional/provokatif, ajakan mendesak segera menyebarkan, ketiadaan rujukan resmi).",
                    "input": "-"
                },
                {
                    "name": "Edukasi interpretasi hasil analisis",
                    "access": "Pengguna",
                    "desc": "Memberikan kartu penjelasan naratif mengenai arti dari skor keyakinan sebagai sinyal waspada awal bagi masyarakat sebelum mempercayai informasi.",
                    "input": "-"
                },
                {
                    "name": "Navigasi aksi lanjutan pemeriksaan",
                    "access": "Pengguna",
                    "desc": "Menyediakan tombol pintas untuk langsung memeriksa teks berita lain secara instan atau kembali ke daftar arsip riwayat.",
                    "input": "Tombol 'Periksa Teks Lain' / 'Kembali ke Riwayat'"
                },
                {
                    "name": "Catatan disclaimer verifikasi fakta",
                    "access": "Pengguna",
                    "desc": "Menyajikan pernyataan batas akurasi statistik model serta rujukan kanal pengecekan fakta independen (Komdigi / CekFakta / TurnBackHoax).",
                    "input": "-"
                }
            ]
        },
        {
            "name": "Riwayat Pemeriksaan (History)",
            "subfeatures": [
                {
                    "name": "Daftar arsip pemeriksaan teks",
                    "access": "Pengguna",
                    "desc": "Menyajikan daftar seluruh riwayat pengujian klaim yang tersimpan di perangkat lengkap dengan waktu WIB, jenis analisis NLP, dan sumber klaim.",
                    "input": "-"
                },
                {
                    "name": "Pencarian riwayat berdasarkan kata kunci",
                    "access": "Pengguna",
                    "desc": "Memfilter dan mencari riwayat pengujian teks berdasarkan kata kunci atau kalimat klaim yang dimasukkan pengguna.",
                    "input": "Input kata kunci pencarian"
                },
                {
                    "name": "Filter kategori status hasil",
                    "access": "Pengguna",
                    "desc": "Menyaring tampilan daftar riwayat berdasarkan kategori klasifikasi (Semua, Terindikasi Hoaks, Valid).",
                    "input": "Pilihan filter chips status"
                },
                {
                    "name": "Halaman detail riwayat tersimpan",
                    "access": "Pengguna",
                    "desc": "Membuka kembali rekaman lengkap pemeriksaan masa lalu termasuk naskah klaim utuh dan tautan ke detail analisis NLP.",
                    "input": "Ketukan pada kartu riwayat"
                },
                {
                    "name": "Penghapusan catatan riwayat",
                    "access": "Pengguna",
                    "desc": "Menghapus satu atau seluruh catatan riwayat simulasi verifikasi teks yang tersimpan di memori perangkat.",
                    "input": "Konfirmasi tombol 'Hapus'"
                }
            ]
        },
        {
            "name": "Profil & Preferensi Pengguna",
            "subfeatures": [
                {
                    "name": "Informasi akun & statistik pengguna",
                    "access": "Pengguna",
                    "desc": "Menampilkan foto profil, nama pengguna, status keanggotaan, serta akumulasi total klaim yang diuji pada bulan berjalan.",
                    "input": "Tombol 'Edit Profil'"
                },
                {
                    "name": "Pengaturan Mode Gelap (Dark Mode)",
                    "access": "Pengguna",
                    "desc": "Mengubah preferensi tema visual antarmuka aplikasi antara tema terang (Light) dan tema gelap (Dark) secara dinamis.",
                    "input": "Switch toggle Mode Gelap"
                },
                {
                    "name": "Menu informasi & kebijakan privasi",
                    "access": "Pengguna",
                    "desc": "Menampilkan lembar informasi profil aplikasi HoaxCheck, nomor versi aplikasi v1.0.0, dan klausul kebijakan privasi data pengguna.",
                    "input": "Pilihan menu navigasi"
                },
                {
                    "name": "Layanan bantuan & lapor masalah",
                    "access": "Pengguna",
                    "desc": "Menyediakan form pengaduan isu kendala aplikasi serta pelaporan klaim berita hoaks baru yang belum dikenali sistem.",
                    "input": "Form teks laporan masalah"
                },
                {
                    "name": "Logout sesi akun pengguna",
                    "access": "Pengguna",
                    "desc": "Mengakhiri sesi masuk pengguna pada perangkat dan mengembalikan antarmuka ke tampilan awal.",
                    "input": "Konfirmasi tombol 'Keluar Akun'"
                }
            ]
        }
    ]
    populate_feature_sheet(ws_mobile, "FITUR", mobile_groups)

    # =========================================================================
    # 3. SHEET: Fitur - Web
    # =========================================================================
    ws_web = wb.create_sheet(title="Fitur - Web")
    web_groups = [
        {
            "name": "Autentikasi Administrator",
            "subfeatures": [
                {
                    "name": "Login admin",
                    "access": "Admin",
                    "desc": "Otentikasi masuk ke web dashboard administratif untuk mengelola dataset dan model verifikasi berita.",
                    "input": "Username dan kata sandi admin"
                },
                {
                    "name": "Logout admin",
                    "access": "Admin",
                    "desc": "Mengakhiri sesi administratif dan menghapus token/sesi web dashboard.",
                    "input": "Tombol 'Logout'"
                }
            ]
        },
        {
            "name": "Dashboard Utama Monitoring",
            "subfeatures": [
                {
                    "name": "Ringkasan metrik model clustering",
                    "access": "Admin",
                    "desc": "Menampilkan evaluasi performa model clustering aktif (Silhouette Score, Davies-Bouldin Index, Calinski-Harabasz Index).",
                    "input": "-"
                },
                {
                    "name": "Monitoring volume & distribusi dataset",
                    "access": "Admin",
                    "desc": "Menyajikan ringkasan jumlah baris dataset berita valid vs berita hoaks Komdigi serta sebaran topik/kategori isu.",
                    "input": "-"
                },
                {
                    "name": "Log aktivitas inferensi API",
                    "access": "Admin",
                    "desc": "Menampilkan grafik dan riwayat jumlah permintaan prediksi/verifikasi teks yang masuk dari aplikasi mobile.",
                    "input": "Filter rentang tanggal"
                }
            ]
        },
        {
            "name": "Manajemen Dataset & Scraper",
            "subfeatures": [
                {
                    "name": "Automated RSS Scraper multi-kanal",
                    "access": "Admin",
                    "desc": "Mengambil artikel berita valid terpercaya secara otomatis melalui RSS feed resmi (Antara News, CNN Indonesia, CNBC Indonesia).",
                    "input": "Pilihan portal berita & batas kuota artikel"
                },
                {
                    "name": "Unggah dataset hoaks / berita baru",
                    "access": "Admin",
                    "desc": "Mengunggah berkas dataset baru (.csv, .xlsx, .xls) ke folder uploads server untuk memperkaya korpus data pelatihan.",
                    "input": "Berkas file dataset (.csv, .xlsx)"
                },
                {
                    "name": "Pratinjau dataset tabular interaktif",
                    "access": "Admin",
                    "desc": "Menampilkan tabel sampel 10 baris pertama data yang berhasil diimpor lengkap dengan nama berkas, ukuran, dan tanggal modifikasi.",
                    "input": "Pilihan berkas pada daftar dataset"
                },
                {
                    "name": "Unduh berkas dataset (Export Data)",
                    "access": "Admin",
                    "desc": "Mengunduh arsip berkas dataset berita valid (e.g. berita_valid_terbaru.csv) atau data komdigi_hoaks ke komputer lokal admin.",
                    "input": "Tombol 'Unduh' pada item berkas"
                },
                {
                    "name": "Penghapusan berkas dataset",
                    "access": "Admin",
                    "desc": "Menghapus file dataset yang sudah tidak digunakan atau usang dari direktori server uploads.",
                    "input": "Tombol 'Hapus' berkas"
                }
            ]
        },
        {
            "name": "Pengelolaan Model ML & NLP",
            "subfeatures": [
                {
                    "name": "Eksekusi pipeline clustering teks",
                    "access": "Admin",
                    "desc": "Menjalankan pipeline praproses teks (tokenisasi, stopword, TF-IDF) dan clustering (K-Means, DBSCAN, HDBSCAN, Hierarchical) secara terprogram.",
                    "input": "Pilihan algoritma dan metrik jarak (Cosine, Euclidean, Manhattan)"
                },
                {
                    "name": "Komparasi performa algoritma",
                    "access": "Admin",
                    "desc": "Membandingkan skor pemisahan klaster antara variasi model clustering dan metrik jarak untuk menentukan konfigurasi model terbaik.",
                    "input": "Parameter uji klaster"
                },
                {
                    "name": "Pembaruan model aktif",
                    "access": "Admin",
                    "desc": "Menyimpan dan menetapkan model clustering/klasifikasi terpilih sebagai inference engine utama (.pkl / .joblib) untuk melayani API.",
                    "input": "Konfirmasi terapkan model"
                }
            ]
        }
    ]
    populate_feature_sheet(ws_web, "FITUR", web_groups)

    # =========================================================================
    # 4. SHEET: Fitur - Cloud
    # =========================================================================
    ws_cloud = wb.create_sheet(title="Fitur - Cloud")
    cloud_groups = [
        {
            "name": "Text Inference Engine API",
            "subfeatures": [
                {
                    "name": "Penerima request inferensi teks",
                    "access": "Sistem (Mobile via API)",
                    "desc": "Menangkap data teks klaim berita melalui metode HTTP POST pada endpoint /api/v1/detect.",
                    "input": "Payload JSON (text_content, user_id, source_type)"
                },
                {
                    "name": "Validasi skema masukan",
                    "access": "Sistem (Backend)",
                    "desc": "Memeriksa kelengkapan format JSON, batasan minimum panjang karakter, dan membersihkan teks masukan dari karakter anomali.",
                    "input": "Format data JSON"
                },
                {
                    "name": "Eksekusi inferensi model NLP",
                    "access": "Sistem (Backend)",
                    "desc": "Memuat berkas serialisasi model (.pkl) dan menghitung estimasi probabilitas status hoaks serta ekstraksi kata kunci.",
                    "input": "Data masukan tervalidasi"
                },
                {
                    "name": "Respons hasil inferensi",
                    "access": "Sistem (Mobile via API)",
                    "desc": "Mengembalikan keluaran analisis status (Terindikasi Hoaks/Valid), confidence score (%), latensi (detik), dan pola indikator ke format JSON.",
                    "input": "-"
                }
            ]
        },
        {
            "name": "Fact-Checking Reference API",
            "subfeatures": [
                {
                    "name": "Penyedia data referensi berita valid",
                    "access": "Sistem (Mobile via API)",
                    "desc": "Mengirimkan artikel berita valid terpercaya dari database sebagai rujukan perbandingan informasi ke aplikasi mobile.",
                    "input": "Query parameter topik / kata kunci"
                },
                {
                    "name": "Penyedia data statistik pasar berita",
                    "access": "Sistem (Mobile via API)",
                    "desc": "Mengirimkan data agregasi total pengujian, rasio hoaks vs valid, dan tren topik disinformasi ke beranda mobile.",
                    "input": "Query parameter periode (opsional)"
                }
            ]
        },
        {
            "name": "Sinkronisasi Riwayat Cloud",
            "subfeatures": [
                {
                    "name": "Sinkronisasi riwayat pemeriksaan",
                    "access": "Sistem (Mobile via API)",
                    "desc": "Mengunggah log pemeriksaan dari perangkat mobile ke server cloud untuk cadangan riwayat pengguna terdaftar.",
                    "input": "Payload JSON riwayat"
                },
                {
                    "name": "Sinkronisasi penghapusan riwayat",
                    "access": "Sistem (Mobile via API)",
                    "desc": "Menghapus log riwayat di server cloud ketika pengguna melakukan penghapusan histori dari perangkat mobile.",
                    "input": "ID pengguna dan ID riwayat"
                }
            ]
        },
        {
            "name": "Infrastruktur Cloud & Keamanan",
            "subfeatures": [
                {
                    "name": "Pengendali CORS & Security",
                    "access": "Sistem (Server Cloud)",
                    "desc": "Mengatur izin akses lintas domain untuk komunikasi aman Flutter mobile dan Web Admin serta otentikasi Bearer Token.",
                    "input": "Konfigurasi header HTTP"
                },
                {
                    "name": "Rate limiting & proteksi request",
                    "access": "Sistem (Server Cloud)",
                    "desc": "Membatasi jumlah request inferensi per IP per menit untuk mencegah penyalahgunaan dan serangan denial of service (DoS).",
                    "input": "Alamat IP pemohon"
                },
                {
                    "name": "Penanganan kode status HTTP",
                    "access": "Sistem (Server Cloud)",
                    "desc": "Mengembalikan kode respons status standar (200 OK, 400 Bad Request, 422 Unprocessable Entity, 429 Too Many Requests, 500 Server Error).",
                    "input": "-"
                }
            ]
        }
    ]
    populate_feature_sheet(ws_cloud, "FITUR", cloud_groups)

    # =========================================================================
    # 5. SHEET: Use Case
    # =========================================================================
    ws_uc = wb.create_sheet(title="Use Case")
    ws_uc.views.sheetView[0].showGridLines = True

    ws_uc.column_dimensions["A"].width = 6.0
    ws_uc.column_dimensions["B"].width = 12.0
    ws_uc.column_dimensions["C"].width = 28.0
    ws_uc.column_dimensions["D"].width = 16.0
    ws_uc.column_dimensions["E"].width = 40.0
    ws_uc.column_dimensions["F"].width = 30.0
    ws_uc.column_dimensions["G"].width = 60.0
    ws_uc.column_dimensions["H"].width = 32.0

    ws_uc.row_dimensions[1].height = 20.0
    ws_uc.row_dimensions[2].height = 24.0

    cell_uc_title = ws_uc.cell(row=1, column=1, value="USE CASE SPECIFICATION")
    cell_uc_title.font = font_title
    cell_uc_title.fill = fill_cyan
    cell_uc_title.alignment = Alignment(horizontal="center", vertical="center")
    cell_uc_title.border = full_thin_border

    for c in range(2, 9):
        c_cell = ws_uc.cell(row=1, column=c, value="")
        c_cell.border = full_thin_border

    uc_headers = ["No", "ID", "Nama Use Case", "Aktor", "Deskripsi Singkat", "Pra-Kondisi", "Alur Utama (Main Flow)", "Pasca-Kondisi"]
    for c_idx, h_text in enumerate(uc_headers, start=1):
        c_cell = ws_uc.cell(row=2, column=c_idx, value=h_text)
        c_cell.font = font_header
        c_cell.fill = fill_cyan
        c_cell.alignment = Alignment(horizontal="center", vertical="center")
        c_cell.border = full_thin_border

    use_cases = [
        (
            1, "UC-01", "Verifikasi Teks Berita Cepat", "Pengguna",
            "Pengguna memeriksa potensi keabsahan klaim berita atau pesan berantai melalui input teks langsung pada aplikasi mobile.",
            "Aplikasi terbuka dan berada pada halaman Deteksi Teks.",
            "1. Pengguna memasukkan atau menempel teks klaim berita ke textfield.\n2. Pengguna menekan tombol 'Periksa Keaslian Teks'.\n3. Sistem mengirimkan teks ke Inference Engine API.\n4. Sistem menampilkan hasil analisis beserta confidence score dan badge klasifikasi.",
            "Hasil analisis tersimpan ke riwayat dan pengguna melihat status verifikasi."
        ),
        (
            2, "UC-02", "Deteksi Otomatis dari Clipboard", "Pengguna",
            "Aplikasi secara cerdas mendeteksi teks yang baru disalin ke clipboard perangkat dan menawarkannya langsung ke pengguna.",
            "Pengguna telah menyalin teks di luar aplikasi (misal dari WhatsApp) dan membuka tab Deteksi.",
            "1. Aplikasi memeriksa clipboard saat halaman dimuat.\n2. Banner saran tempel muncul jika terdapat teks baru.\n3. Pengguna menekan tombol 'Gunakan'.\n4. Teks otomatis tertempel ke dalam textfield deteksi.",
            "Form deteksi terisi teks klaim tanpa perlu penempelan manual."
        ),
        (
            3, "UC-03", "Meninjau Detail Indikator NLP", "Pengguna",
            "Pengguna meninjau rincian indikator linguistik manipulatif yang ditemukan pada teks yang diuji.",
            "Pengguna telah melakukan pemeriksaan teks atau membuka riwayat pemeriksaan.",
            "1. Pengguna berada di halaman Detail Analisis.\n2. Sistem menyajikan Circular Confidence Dial (%).\n3. Pengguna meninjau daftar kartu pola bahasa manipulatif (provokatif, ajakan sebar, tanpa rujukan).\n4. Pengguna membaca edukasi interpretasi hasil.",
            "Pengguna memahami alasan di balik kesimpulan klasifikasi model."
        ),
        (
            4, "UC-04", "Menelusuri & Memfilter Riwayat", "Pengguna",
            "Pengguna mencari atau memfilter log rekaman pemeriksaan teks yang pernah dilakukan sebelumnya.",
            "Terdapat minimal satu riwayat pengujian tersimpan.",
            "1. Pengguna membuka tab Riwayat.\n2. Pengguna memasukkan kata kunci pencarian atau memilih chip filter (Semua, Terindikasi Hoaks, Valid).\n3. Sistem memperbarui daftar kartu riwayat sesuai filter.",
            "Daftar riwayat tersaring sesuai kriteria pencarian pengguna."
        ),
        (
            5, "UC-05", "Mengubah Tema (Dark/Light Mode)", "Pengguna",
            "Pengguna menyesuaikan preferensi tema visual aplikasi mobile sesuai kenyamanan visual.",
            "Pengguna berada di halaman Profil & Preferensi.",
            "1. Pengguna melihat toggle switch 'Mode Gelap'.\n2. Pengguna mengaktifkan atau menonaktifkan sakelar.\n3. Sistem secara dinamis mengubah tema warna antarmuka aplikasi secara instan.",
            "Preferensi tema visual tersimpan di memori perangkat."
        ),
        (
            6, "UC-06", "Otentikasi Administrator", "Admin",
            "Administrator melakukan login ke sistem Web Admin untuk mengelola dataset dan model inferensi.",
            "Web dashboard admin terbuka pada browser.",
            "1. Admin memasukkan username dan kata sandi.\n2. Sistem memvalidasi kredensial admin.\n3. Sistem membuat sesi administratif dan mengarahkan ke dashboard.",
            "Admin berhasil masuk ke dashboard utama pengelolaan."
        ),
        (
            7, "UC-07", "Scraping Otomatis Berita Valid", "Admin",
            "Administrator mengeksekusi scraper RSS feed untuk mengunduh berita valid dari media resmi terpercaya.",
            "Admin telah login dan berada di tab Scraper Berita.",
            "1. Admin memilih portal media (Antara, CNN, CNBC, Semua).\n2. Admin menentukan kuota batas artikel.\n3. Admin menekan tombol 'Mulai Scrape Berita'.\n4. Sistem mengambil RSS feeds dan menyimpannya sebagai file CSV.",
            "File CSV berita valid terbaru tersimpan di folder server uploads."
        ),
        (
            8, "UC-08", "Pengelolaan & Pratinjau Dataset", "Admin",
            "Administrator mengunggah berkas data baru, melihat pratinjau 10 baris pertama, atau mengunduh dataset.",
            "Admin berada di tab Manajemen Dataset pada Web Admin.",
            "1. Admin memilih file dataset (.csv, .xlsx) dan menekan 'Unggah'.\n2. Sistem menyimpan dan memverifikasi berkas.\n3. Admin dapat menekan tombol 'Pratinjau' untuk melihat tabel 10 baris pertama atau 'Unduh'.",
            "Korpus dataset diperbarui dan siap digunakan untuk pelatihan model."
        ),
        (
            9, "UC-09", "Eksperimen Model Clustering ML", "Admin",
            "Administrator menjalankan pelatihan ulang klasterisasi topik hoaks dan membandingkan performa model.",
            "Dataset hoaks Komdigi dan berita valid tersedia di server.",
            "1. Admin memilih algoritma (K-Means / HDBSCAN / Hierarchical) dan metrik jarak.\n2. Sistem mengeksekusi pipeline TF-IDF dan clustering.\n3. Sistem menampilkan metrik Silhouette, Davies-Bouldin, dan Calinski-Harabasz.\n4. Admin menerapkan model terbaik sebagai model aktif.",
            "Model inferensi terbarui (.pkl) dan aktif melayani API mobile."
        ),
        (
            10, "UC-10", "Eksekusi Inferensi API Real-Time", "Sistem (Mobile API)",
            "Aplikasi mobile memanggil endpoint REST API untuk meminta analisis teks berita dari inference engine.",
            "Server backend cloud aktif dan menerima HTTP POST request.",
            "1. Mobile mengirim POST /api/v1/detect dengan payload teks.\n2. Backend memvalidasi masukan dan menjalankan model NLP.\n3. Backend menghitung probabilitas status dan indikator.\n4. Backend mengembalikan response JSON status 200 OK.",
            "Mobile menerima payload hasil prediksi dan menampilkannya ke layar."
        )
    ]

    for uc_idx, uc in enumerate(use_cases, start=3):
        color_hex = palette[uc_idx % len(palette)]
        uc_fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
        ws_uc.row_dimensions[uc_idx].height = 42.0

        for col_idx, val in enumerate(uc, start=1):
            cell = ws_uc.cell(row=uc_idx, column=col_idx, value=val)
            cell.font = font_regular
            cell.fill = uc_fill
            cell.border = full_thin_border
            if col_idx in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # =========================================================================
    # 6. SHEET: Schema Database
    # =========================================================================
    ws_schema = wb.create_sheet(title="Schema Database")
    ws_schema.views.sheetView[0].showGridLines = True

    ws_schema.column_dimensions["A"].width = 24.0
    ws_schema.column_dimensions["B"].width = 28.0
    ws_schema.column_dimensions["C"].width = 42.0
    ws_schema.column_dimensions["D"].width = 48.0
    ws_schema.column_dimensions["E"].width = 38.0

    ws_schema.row_dimensions[1].height = 20.0
    ws_schema.row_dimensions[2].height = 24.0

    cell_sch_title = ws_schema.cell(row=1, column=1, value="Schema Database")
    cell_sch_title.font = font_title
    cell_sch_title.fill = fill_schema_header
    cell_sch_title.alignment = Alignment(horizontal="left", vertical="center")

    sch_headers = ["Table Name / Field Name", "Field Type", "Validation / Constraint", "Deskripsi", "Sample"]
    for c_idx, h_text in enumerate(sch_headers, start=1):
        cell = ws_schema.cell(row=2, column=c_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_schema_header
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = full_thin_border

    tables = [
        {
            "name": "admins",
            "fields": [
                ("id", "UUID", "PK, unique, default uuid_generate_v4()", "Primary key identitas unik akun administrator", "a0eebc99-9c0b-4ef8-bb6d-6bb9bd380a11"),
                ("username", "VARCHAR(50)", "unique, required", "Nama pengguna untuk otentikasi web dashboard admin", "admin_hoaxcheck"),
                ("email", "VARCHAR(100)", "unique, required, format email", "Alamat email resmi administrator pengelola", "admin@hoaxcheck.id"),
                ("password_hash", "VARCHAR(255)", "required", "Hash kata sandi terenkripsi menggunakan bcrypt", "$2b$12$e8m4XpQ9..."),
                ("full_name", "VARCHAR(100)", "required", "Nama lengkap administrator pengelola sistem", "Muhammad Rayhan"),
                ("role", "VARCHAR(30)", "default 'SUPER_ADMIN'", "Peran hak akses administratif (SUPER_ADMIN, ML_ENGINEER)", "SUPER_ADMIN"),
                ("is_active", "BOOLEAN", "default TRUE", "Status keaktifan akun administrator", True),
                ("created_at", "TIMESTAMP WITH TIME ZONE", "autogenerate, default CURRENT_TIMESTAMP", "Waktu pembuatan akun administrator", "2026-10-01 08:30:00+07"),
                ("updated_at", "TIMESTAMP WITH TIME ZONE", "autogenerate, default CURRENT_TIMESTAMP", "Waktu pembaruan profil administrator", "2026-10-01 08:30:00+07")
            ]
        },
        {
            "name": "dataset_metadata",
            "fields": [
                ("id", "UUID", "PK, unique, default uuid_generate_v4()", "Primary key riwayat berkas dataset fisik di server", "b1ffcd88-8d1c-4ef8-bb6d-7cc0ce491b22"),
                ("file_name", "VARCHAR(255)", "required", "Nama berkas fisik dataset yang diunggah / di-scrape", "komdigi_hoaks.csv"),
                ("file_path", "VARCHAR(500)", "required", "Lokasi penyimpanan fisik file di storage server", "/uploads/komdigi_hoaks.csv"),
                ("dataset_type", "VARCHAR(50)", "required", "Kategori dataset (HOAX_KOMDIGI, VALID_RSS, COMBINED)", "HOAX_KOMDIGI"),
                ("total_records", "INTEGER", "required, min 0", "Jumlah total baris data di dalam berkas dataset", 4850),
                ("file_size_bytes", "BIGINT", "required", "Ukuran berkas dalam satuan byte", 1245800),
                ("source_portal", "VARCHAR(100)", "optional", "Portal asal dataset (Komdigi / Antara / CNN / CNBC)", "Kominfo / Komdigi"),
                ("uploaded_by", "UUID", "FK optional, references admins(id)", "Relasi admin yang mengunggah atau mengeksekusi scraper", "a0eebc99-9c0b-4ef8-bb6d-6bb9bd380a11"),
                ("created_at", "TIMESTAMP WITH TIME ZONE", "autogenerate, default CURRENT_TIMESTAMP", "Waktu berkas dataset ditambahkan ke server", "2026-10-02 09:15:00+07")
            ]
        },
        {
            "name": "scraped_articles",
            "fields": [
                ("id", "BIGSERIAL", "PK, unique, autoincrement", "Primary key identitas artikel hasil scraping RSS resmi", 1),
                ("source", "VARCHAR(100)", "required", "Nama media resmi penerbit berita (Antara, CNN, CNBC)", "Antara News (Kantor Berita Negara)"),
                ("title", "TEXT", "required", "Judul naskah artikel berita valid terpercaya", "Pemerintah Klarifikasi Hoaks Informasi Pungutan Biaya Booster"),
                ("slug", "VARCHAR(255)", "optional", "Slug pengenal URL artikel berita", "pemerintah-klarifikasi-hoaks-booster"),
                ("excerpt", "TEXT", "optional", "Ringkasan naskah artikel berita valid", "Kementerian Kesehatan menegaskan bahwa program vaksinasi tetap gratis..."),
                ("url", "VARCHAR(1000)", "unique, required", "Alamat tautan permanen berita resmi daring", "https://www.antaranews.com/berita/12345/klarifikasi-booster"),
                ("published_at", "VARCHAR(100)", "optional", "Waktu publikasi resmi yang tertera pada RSS feed", "Wed, 14 Oct 2026 08:30:00 GMT"),
                ("label", "VARCHAR(20)", "default 'valid'", "Label klasifikasi artikel untuk referensi korpus", "valid"),
                ("created_at", "TIMESTAMP WITH TIME ZONE", "autogenerate, default CURRENT_TIMESTAMP", "Waktu artikel disimpan ke basis data", "2026-10-14 08:45:00+07")
            ]
        },
        {
            "name": "ml_models",
            "fields": [
                ("id", "UUID", "PK, unique, default uuid_generate_v4()", "Primary key entri model machine learning / NLP", "c2eedf77-7e2b-4ef8-bb6d-8dd1df502c33"),
                ("model_name", "VARCHAR(100)", "required", "Nama algoritma / arsitektur model NLP", "HDBSCAN Text Clustering + IndoBERT"),
                ("version", "VARCHAR(20)", "required", "Nomor versi rilis model machine learning", "v1.2.0"),
                ("model_file_path", "VARCHAR(500)", "required", "Path berkas serialisasi (.pkl / .joblib) di server", "/models/hdbscan_cosine_v1.pkl"),
                ("distance_metric", "VARCHAR(50)", "required", "Metrik jarak perhitungan klaster (Cosine, Euclidean, Manhattan)", "Cosine Similarity"),
                ("feature_extractor", "VARCHAR(100)", "required", "Metode ekstraksi representasi teks", "TF-IDF (max_features=5000)"),
                ("silhouette_score", "NUMERIC(6,4)", "optional, range -1.0 s/d 1.0", "Skor evaluasi kohesi dan separasi klaster", 0.6842),
                ("davies_bouldin_index", "NUMERIC(6,4)", "optional, non-negative", "Skor indeks evaluasi pemisahan Davies-Bouldin", 0.7421),
                ("calinski_harabasz_index", "NUMERIC(10,2)", "optional, non-negative", "Skor indeks rasio dispersi Calinski-Harabasz", 1420.50),
                ("accuracy_f1", "NUMERIC(6,4)", "optional, range 0.0 s/d 1.0", "Skor evaluasi F1-Score klasifikasi probabilitas", 0.9150),
                ("is_active", "BOOLEAN", "default FALSE, unique index jika TRUE", "Menandakan model utama yang melayani API inferensi", True),
                ("trained_by", "UUID", "FK required, references admins(id)", "Relasi admin yang mengeksekusi pelatihan model", "a0eebc99-9c0b-4ef8-bb6d-6bb9bd380a11"),
                ("created_at", "TIMESTAMP WITH TIME ZONE", "autogenerate, default CURRENT_TIMESTAMP", "Waktu model selesai dievaluasi dan dicatat", "2026-10-05 14:20:00+07")
            ]
        },
        {
            "name": "detection_logs",
            "fields": [
                ("id", "UUID", "PK, unique, default uuid_generate_v4()", "Primary key log permintaan inferensi prediksi dari mobile", "d3ffea66-6f3a-4ef8-bb6d-9ee2ea613d44"),
                ("analysis_code", "VARCHAR(20)", "required, unique", "Kode identitas unik analisis ramah pengguna", "#HC-98234"),
                ("text_input", "TEXT", "required", "Teks narasi klaim berita yang diuji oleh pengguna", "Beredar pesan berantai mengklaim vaksinasi booster gratis dibatalkan..."),
                ("text_char_count", "INTEGER", "required, min 1", "Jumlah karakter pada teks masukan pengujian", 245),
                ("predicted_label", "VARCHAR(50)", "required", "Hasil klasifikasi akhir inferensi teks", "Terindikasi Hoaks"),
                ("confidence_score", "NUMERIC(5,2)", "required, range 0 s/d 100", "Skor persentase keyakinan model terhadap kesimpulan (%)", 87.00),
                ("is_hoax", "BOOLEAN", "required", "Flag status biner (True jika hoaks, False jika valid)", True),
                ("detected_patterns_count", "INTEGER", "default 0", "Jumlah pola indikator kebahasaan manipulatif yang ditemukan", 3),
                ("model_id", "UUID", "FK required, references ml_models(id)", "Relasi versi model ML yang mengeksekusi inferensi", "c2eedf77-7e2b-4ef8-bb6d-8dd1df502c33"),
                ("latency_seconds", "NUMERIC(5,2)", "optional, non-negative", "Waktu komputasi inferensi dalam satuan detik", 1.20),
                ("source_channel", "VARCHAR(50)", "optional", "Sumber asal masukan teks klaim berita", "WhatsApp Forward"),
                ("client_ip", "VARCHAR(45)", "optional", "Alamat IP perangkat pemohon (IPv4/IPv6)", "182.253.110.45"),
                ("created_at", "TIMESTAMP WITH TIME ZONE", "autogenerate, default CURRENT_TIMESTAMP", "Waktu permintaan deteksi dieksekusi", "2026-10-14 10:42:15+07")
            ]
        },
        {
            "name": "linguistic_indicators",
            "fields": [
                ("id", "BIGSERIAL", "PK, unique, autoincrement", "Primary key rincian temuan pola linguistik per pengujian", 1),
                ("detection_log_id", "UUID", "FK required, references detection_logs(id)", "Relasi log permintaan inferensi deteksi teks", "d3ffea66-6f3a-4ef8-bb6d-9ee2ea613d44"),
                ("indicator_name", "VARCHAR(100)", "required", "Nama pola linguistik yang terdeteksi pada teks", "Bahasa Sensasional & Provokatif"),
                ("weight_level", "VARCHAR(20)", "required", "Bobot pengaruh indikator (Bobot Tinggi, Bobot Sedang, Bobot Rendah)", "Bobot Tinggi"),
                ("indicator_description", "TEXT", "required", "Rincian penjelasan temuan karakteristik kebahasaan", "Teks menggunakan ungkapan hiperbolis yang memancing kepanikan emosional."),
                ("created_at", "TIMESTAMP WITH TIME ZONE", "autogenerate, default CURRENT_TIMESTAMP", "Waktu indikator tercatat dalam sistem", "2026-10-14 10:42:15+07")
            ]
        },
        {
            "name": "user_histories",
            "fields": [
                ("id", "UUID", "PK, unique, default uuid_generate_v4()", "Primary key catatan riwayat pemeriksaan di perangkat mobile", "f5bbea11-1a2b-4ef8-bb6d-1111aa223344"),
                ("user_device_id", "VARCHAR(100)", "required", "Pengenal unik instalasi perangkat aplikasi mobile", "device_usr_882910"),
                ("title", "VARCHAR(255)", "required", "Judul ringkas naskah klaim yang diuji", "Pesan Berantai Klaim Vaksinasi Booster Mandiri"),
                ("snippet", "TEXT", "required", "Cuplikan teks klaim yang diuji pengguna", "Beredar pesan berantai mengklaim vaksinasi booster gratis dibatalkan..."),
                ("classification", "VARCHAR(50)", "required", "Status hasil klasifikasi (Terindikasi Hoaks / Valid)", "Terindikasi Hoaks"),
                ("confidence_str", "VARCHAR(10)", "required", "Teks persentase keyakinan model", "87%"),
                ("is_hoax", "BOOLEAN", "required", "Penanda status apakah klaim berupa hoaks", True),
                ("analysis_type", "VARCHAR(50)", "default 'Analisis NLP'", "Jenis metode analisis inferensi yang digunakan", "Analisis NLP"),
                ("source", "VARCHAR(50)", "optional", "Sumber klaim berita yang diuji", "WhatsApp Forward"),
                ("tested_at", "TIMESTAMP WITH TIME ZONE", "autogenerate, default CURRENT_TIMESTAMP", "Waktu pengujian klaim dilakukan oleh pengguna", "2026-10-14 10:42:00+07")
            ]
        },
        {
            "name": "system_audit_logs",
            "fields": [
                ("id", "UUID", "PK, unique, default uuid_generate_v4()", "Primary key pencatatan audit aktivitas admin", "e4aaeb55-5f2b-4ef8-bb6d-0ff3fa724e55"),
                ("admin_id", "UUID", "FK required, references admins(id)", "Relasi administrator yang melakukan aksi perubahan", "a0eebc99-9c0b-4ef8-bb6d-6bb9bd380a11"),
                ("action_type", "VARCHAR(50)", "required", "Jenis operasi (UPLOAD_DATASET, RSS_SCRAPE, RETRAIN_MODEL, SWITCH_MODEL, DELETE)", "RSS_SCRAPE"),
                ("module_name", "VARCHAR(50)", "required", "Modul sistem tempat aksi dilakukan (DATASET_MANAGER, SCRAPER, ML_REGISTRY)", "SCRAPER"),
                ("description", "TEXT", "required", "Rincian deskripsi tindakan yang dijalankan admin", "Berhasil scraping 150 artikel berita valid dari Antara News dan CNBC Indonesia"),
                ("ip_address", "VARCHAR(45)", "optional", "Alamat IP admin saat melakukan aksi", "192.168.1.18"),
                ("user_agent", "TEXT", "optional", "Informasi peramban dan perangkat sistem operasi admin", "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/129.0.0.0"),
                ("created_at", "TIMESTAMP WITH TIME ZONE", "autogenerate, default CURRENT_TIMESTAMP", "Waktu kejadian aktivitas tercatat di sistem", "2026-10-14 09:30:00+07")
            ]
        }
    ]

    curr_row = 4
    for tbl in tables:
        # Table Header
        ws_schema.row_dimensions[curr_row].height = 22.0
        tbl_cell = ws_schema.cell(row=curr_row, column=1, value=tbl["name"])
        tbl_cell.font = font_table_name
        tbl_cell.fill = fill_schema_table
        tbl_cell.alignment = Alignment(horizontal="left", vertical="center")
        tbl_cell.border = full_thin_border

        for c in range(2, 6):
            c_cell = ws_schema.cell(row=curr_row, column=c, value="")
            c_cell.fill = fill_schema_table
            c_cell.border = full_thin_border

        curr_row += 1

        # Fields
        for fld in tbl["fields"]:
            ws_schema.row_dimensions[curr_row].height = 20.0
            for c_idx, val in enumerate(fld, start=1):
                cell = ws_schema.cell(row=curr_row, column=c_idx, value=val)
                cell.font = font_regular
                cell.border = full_thin_border
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            curr_row += 1

        # Blank row spacer between tables
        curr_row += 1

    # Save to target destinations
    target_paths = [
        r"c:\Users\azure\OneDrive\Documents\PBL_SEM5\HoaxCheckML\Dokumen Perencanaan - HoaxCheck.xlsx",
        r"C:\Users\azure\OneDrive\Documents\Mobile Dev\HoaxCheck\Dokumen Perencanaan - HoaxCheck.xlsx",
        r"C:\Users\azure\Downloads\Dokumen Perencanaan - HoaxCheck.xlsx"
    ]

    for p in target_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        wb.save(p)
        print(f"Successfully generated: {p}")

if __name__ == "__main__":
    create_planning_document()
